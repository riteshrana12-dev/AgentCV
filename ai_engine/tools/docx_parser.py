import io
from typing import Iterator
from docx import Document as DocxDocument
from langchain_community.document_loaders.blob_loaders import Blob
from langchain_core.document_loaders import BaseBlobParser
from langchain_core.documents import Document as LCDocument


class InMemoryDocxParser(BaseBlobParser):

    def lazy_parse(self, blob: Blob) -> Iterator[LCDocument]:
        """
        Parses a Blob containing DOCX byte data into a LangChain Document.
        """
        try:
            # Wrap in-memory byte stream
            file_stream = io.BytesIO(blob.as_bytes())
            doc = DocxDocument(file_stream)
            
            # Extract main paragraph text
            paragraphs_text = [p.text for p in doc.paragraphs if p.text.strip()]
            
            # Extract text embedded inside tables (frequently used in ATS resumes)
            table_text = []
            for table in doc.tables:
                for row in table.rows:
                    row_data = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_data:
                        table_text.append(" | ".join(row_data))

            # Combine all text content with double newlines to retain paragraph separation
            full_text = "\n\n".join(paragraphs_text + table_text)
            
            # Metadata propagation
            metadata = blob.metadata if blob.metadata else {}
            
            yield LCDocument(page_content=full_text, metadata=metadata)
            
        except Exception as e:
            raise ValueError(f"Failed to parse DOCX file '{blob.source}': {str(e)}") from e
